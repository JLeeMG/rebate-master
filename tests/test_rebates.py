"""The rebate master: four-eyes, history, per-brand approvers, and the workbook reader."""

from datetime import date, datetime
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import select, text
from sqlalchemy.exc import DBAPIError, IntegrityError

from mgrm.auth.roles import Role
from mgrm.data.errors import AlreadyLoaded, LoadRejected
from mgrm.data.rebate_workbooks import load_workbook, parse_tab, propose_rates
from mgrm.domain.entities import Entity
from mgrm.domain.periods import add_months
from mgrm.models import (
    AccrualMode,
    RateType,
    RebateAgreement,
    RebateApproverScope,
    RebateBasis,
    RebateRate,
    ReviewStatus,
)
from mgrm.rebates.service import (
    RebateError,
    Upload,
    agreements_with_rates,
    approve,
    expiry_and_renewal,
    propose_rate,
    reject,
)

from tests.conftest import PDF_BYTES

AS_AT = date(2026, 8, 1)


@pytest.fixture
def people(make_user):
    return {
        "editor": make_user(Role.REBATE_EDITOR, email="editor@example.com"),
        "ken": make_user(Role.REBATE_REVIEWER, email="ken@example.com"),
        "gpm": make_user(Role.BRAND_APPROVER, email="gpm@example.com"),
        "admin": make_user(Role.ADMIN, email="cfo@example.com"),
        "reader": make_user(Role.VIEWER, email="raymond@example.com"),
    }


@pytest.fixture
def agreement(db, people):
    a = RebateAgreement(
        code="MGNZ-NLG-EUFY-1", entity_code="MGNZ", customer_label="Noel Leeming", customer_group_code="NL",
        brand_code="EUFY", product_scope="EUFY (Security & Baby)", rate_type=RateType.REBATE,
        basis=RebateBasis.REBATE_ELIGIBLE_SALES, source_reference="Trading terms 2026", created_by_id=people["editor"].id,
    )
    db.add(a)
    db.flush()
    return a


def propose(db, people, agreement, rate="0.253", start=date(2026, 8, 1), end=None, who="editor"):
    return propose_rate(db, actor=people[who], agreement=agreement, rate=Decimal(rate), effective_from=start,
                        effective_to=end, source_reference="NLG trading terms Aug 2026", reason="new range",
                        evidence=[Upload("NLG email.pdf", PDF_BYTES)])


def refused(db, action, match):
    with pytest.raises((DBAPIError, IntegrityError), match=match):
        with db.begin_nested():
            action()
            db.flush()


# ---------------------------------------------------------------- four-eyes and history


def test_nobody_approves_their_own_rate(db, people, agreement):
    rate = propose(db, people, agreement, who="admin")
    with pytest.raises(RebateError, match="someone else"):
        approve(db, actor=people["admin"], rate=rate)


def test_the_database_refuses_self_approval_too(db, people, agreement):
    rate = propose(db, people, agreement)

    def self_approve():
        db.execute(text("UPDATE rebate_rate SET status='approved', reviewed_by_id=entered_by_id, reviewed_at=now() WHERE id=:i"), {"i": rate.id})

    refused(db, self_approve, "four_eyes")


def test_a_source_document_is_required(db, people, agreement):
    with pytest.raises(RebateError, match="Name the source"):
        propose_rate(db, actor=people["editor"], agreement=agreement, rate=Decimal("0.2"), effective_from=AS_AT,
                     effective_to=None, source_reference=" ", reason="renewal", evidence=[Upload("a.pdf", PDF_BYTES)])


def test_a_rebates_only_user_cannot_enter_or_approve(db, people, agreement):
    with pytest.raises(RebateError):
        propose(db, people, agreement, who="reader")
    rate = propose(db, people, agreement)
    with pytest.raises(RebateError, match="does not approve"):
        approve(db, actor=people["reader"], rate=rate)


def test_approving_a_change_closes_the_old_rate_and_keeps_it(db, people, agreement):
    first = propose(db, people, agreement, rate="0.20", start=date(2025, 4, 1))
    approve(db, actor=people["ken"], rate=first)
    second = propose(db, people, agreement, rate="0.253", start=date(2026, 8, 1))
    approve(db, actor=people["ken"], rate=second)
    db.refresh(first)
    assert first.effective_to == date(2026, 7, 31) and first.rate == Decimal("0.2")
    on = {c.agreement.code: c.rate.rate for c in agreements_with_rates(db, date(2026, 7, 15))}
    assert on["MGNZ-NLG-EUFY-1"] == Decimal("0.2")
    assert agreements_with_rates(db, date(2026, 8, 15))[0].rate.rate == Decimal("0.253")


def test_an_approved_rate_is_history(db, people, agreement):
    rate = propose(db, people, agreement)
    approve(db, actor=people["ken"], rate=rate)
    refused(db, lambda: db.execute(text("UPDATE rebate_rate SET rate = 0.3 WHERE id=:i"), {"i": rate.id}), "cannot be changed")
    refused(db, lambda: db.execute(text("DELETE FROM rebate_rate WHERE id=:i"), {"i": rate.id}), "cannot be deleted")
    db.execute(text("UPDATE rebate_rate SET effective_to = '2026-12-31' WHERE id=:i"), {"i": rate.id})
    refused(db, lambda: db.execute(text("UPDATE rebate_rate SET effective_to = '2027-03-31' WHERE id=:i"), {"i": rate.id}), "only once")


def test_overlapping_approved_rates_are_refused(db, people, agreement):
    approve(db, actor=people["ken"], rate=propose(db, people, agreement, rate="0.2", start=date(2026, 1, 1), end=date(2026, 12, 31)))
    clash = propose(db, people, agreement, rate="0.25", start=date(2026, 6, 1))
    with pytest.raises(RebateError, match="already covers"):
        approve(db, actor=people["ken"], rate=clash)


def test_reject_needs_a_reason(db, people, agreement):
    rate = propose(db, people, agreement)
    with pytest.raises(RebateError, match="why"):
        reject(db, actor=people["ken"], rate=rate, note="")
    reject(db, actor=people["ken"], rate=rate, note="Terms say 25%, not 25.3%")
    assert rate.status is ReviewStatus.REJECTED


def test_brand_approvers_take_over_their_brand(db, people, agreement):
    # Ken approves now. Once EUFY has an assigned approver (its GPM), only they or the administrator may.
    rate = propose(db, people, agreement)
    with pytest.raises(RebateError, match="not an assigned approver"):
        approve(db, actor=people["gpm"], rate=rate)
    db.add(RebateApproverScope(user_id=people["gpm"].id, brand_code="EUFY"))
    db.flush()
    with pytest.raises(RebateError, match="assigned approvers"):
        approve(db, actor=people["ken"], rate=rate)
    approve(db, actor=people["gpm"], rate=rate)
    assert rate.status is ReviewStatus.APPROVED and rate.reviewed_by_id == people["gpm"].id


def test_expiry_and_renewal(db, people, agreement):
    approve(db, actor=people["ken"], rate=propose(db, people, agreement, rate="0.2", start=date(2026, 1, 1), end=date(2026, 10, 31)))
    ending, open_ended = expiry_and_renewal(db, date(2026, 9, 26))
    assert [e.agreement.code for e in ending] == ["MGNZ-NLG-EUFY-1"] and open_ended == []


# ---------------------------------------------------------------- the workbook reader

MONTHS = [add_months(date(2024, 1, 1), i) for i in range(32)]  # Jan 2024 .. Aug 2026


def tab_rows(*, customer="JB Hi Fi", scope1="ALL", rate1=0.1, sales1=None, due1=None, scope2=None, rate2=None,
             sales2=None, due2=None, notes=(), month_rates=None):
    """A grid laid out like the real template (rows 1-4 header, row 7 months, rows 8-14 figures)."""
    width = 2 + len(MONTHS)
    rows = [[None] * width for _ in range(24)]
    rows[0][0] = "Home"
    for i, note in enumerate(notes):
        rows[0][2 + i] = note
    rows[2][:4] = [customer, "Product Groups", scope1, rate1]
    rows[3][:4] = [None, "Product Groups", scope2, rate2]
    for j, month in enumerate(MONTHS):
        rows[6][2 + j] = datetime(month.year, month.month, 1)
        if month_rates and month in month_rates:
            rows[2][2 + j] = month_rates[month]
    rows[7][:2] = ["Sales Ex GST $", scope1]
    rows[8][:2] = ["Sales Ex GST $", scope2]
    rows[9][0] = "Total Sales"
    rows[11][:2] = ["Rebate Due ", rate1]
    rows[12][:2] = ["Rebate Due ", rate2]
    rows[13][0] = "Total Rebate Due"
    for row, series in ((7, sales1), (11, due1), (8, sales2), (12, due2)):
        for month, value in (series or {}).items():
            rows[row][2 + MONTHS.index(month)] = value
    return rows


def constant(rate, months=MONTHS, sales=1000.0):
    return {m: sales for m in months}, {m: round(sales * rate, 2) for m in months}


def proposals(rows, sheet="JB Hi Fi"):
    tab = parse_tab(rows, sheet, Entity.MGAU)
    return tab, [propose_rates(tab, line, AS_AT) for line in tab.lines]


def test_a_constant_rate_becomes_one_open_period():
    sales, due = constant(0.19)
    tab, [[p]] = proposals(tab_rows(rate1=0.19, sales1=sales, due1=due))
    assert (p.rate, p.effective_from, p.effective_to, p.flags) == (Decimal("0.19"), date(2024, 1, 1), None, ())
    assert tab.group_code == "JBH" and tab.rate_type is RateType.REBATE


def test_rate_history_is_rebuilt_and_a_stale_rate_cell_is_flagged():
    # The OfficeWorks pattern: rates written above each month; the column-B cell never updated.
    early, late = MONTHS[:20], MONTHS[20:]
    sales = {m: 1000.0 for m in MONTHS}
    due = {**{m: 145.0 for m in early}, **{m: 190.0 for m in late}}
    rows = tab_rows(customer="OfficeWorks", rate1=0.1, sales1=sales, due1=due,
                    month_rates={**{m: 0.145 for m in early}, **{m: 0.19 for m in late}})
    _, [[old, new]] = proposals(rows, "OfficeWorks")
    assert (old.rate, old.effective_to) == (Decimal("0.145"), date(2025, 8, 31))
    assert (new.rate, new.effective_from, new.effective_to) == (Decimal("0.19"), date(2025, 9, 1), None)
    assert any("says 10%" in f and "19% is proposed" in f for f in new.flags)


def test_months_without_sales_and_cent_rounding_do_not_split_a_period():
    sales = {m: 333.33 for m in MONTHS if m.month != 3}
    due = {m: round(333.33 * 0.01, 2) for m in sales}  # 3.33: rounding puts the applied rate at 0.999%
    _, [[p]] = proposals(tab_rows(rate1=0.01, sales1=sales, due1=due))
    assert p.rate == Decimal("0.01") and p.effective_to is None


def test_impossible_and_unwritten_rates_are_flagged_not_proposed():
    sales, due = constant(0.05)
    due[MONTHS[5]] = sales[MONTHS[5]] * 2.4896  # the MGNZ JB Hi Fi Airport month
    due[MONTHS[9]] = sales[MONTHS[9]] * 0.0366  # a blended or manual month
    tab, [[p]] = proposals(tab_rows(rate1=0.05, sales1=sales, due1=due))
    flags = " | ".join(tab.lines[0].flags)
    assert "Impossible applied rate" in flags and "not written on the tab" in flags
    assert p.rate == Decimal("0.05") and p.effective_to is None


def test_an_agreement_without_recent_sales_is_proposed_as_ended():
    sales, due = constant(0.05, MONTHS[:10])
    _, [[p]] = proposals(tab_rows(rate1=0.05, sales1=sales, due1=due))
    assert p.effective_to == date(2024, 10, 31) and any("No sales since" in f for f in p.flags)


def test_a_new_agreement_takes_its_start_from_its_own_line():
    sales, due = constant(0.15)
    rows = tab_rows(customer="The Good Guys", rate1=0.15, sales1=sales, due1=due,
                    scope2="EUFY 08/26 Onwards", rate2=0.1, notes=("Mova 22% 05/2024 onwards",))
    _, [_, [p]] = proposals(rows, "The Good Guys")
    assert (p.rate, p.effective_from) == (Decimal("0.1"), date(2026, 8, 1))


def test_check_only_needs_the_explicit_phrase():
    sales, due = constant(0.028)
    tab = parse_tab(tab_rows(rate1=0.028, sales1=sales, due1=due, notes=("DO NOT ACCRUE PURELY FOR CHECK",)), "NLG Honor 2.8%", Entity.MGNZ)
    assert tab.accrual_mode is AccrualMode.CHECK_ONLY
    costco = parse_tab(tab_rows(customer="Costco", rate1=0.055, sales1=sales, due1=due,
                                notes=("ensure that Costco do not claim twice do not accrue for Insta360 Sales",)), "Costco", Entity.MGNZ)
    assert costco.accrual_mode is AccrualMode.ACCRUE
    assert any("check whether part of the sales must be excluded" in f for f in costco.flags)
    assert any("No customer group for 'Costco'" in f for f in costco.flags)


def test_excel_float_noise_is_removed():
    sales, due = constant(0.173)
    _, [[p]] = proposals(tab_rows(rate1=0.17300000000000001, sales1=sales, due1=due))
    assert str(p.rate) == "0.173"


def test_a_sheet_that_is_not_an_agreement_is_skipped():
    assert "not an agreement tab" in parse_tab([["Financial Summary"]] * 10, "Financial Summary", Entity.MGAU)


def test_loading_a_workbook(db, people, tmp_path):
    path = tmp_path / "MGNZ 2608 Aug 26 Rebates Reconciliation.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Financial Summary"
    for sheet, rows in (("JB Hi Fi", tab_rows(rate1=0.2, sales1=constant(0.2)[0], due1=constant(0.2)[1])),
                        ("Costco", tab_rows(customer="Costco", rate1=0.055, sales1=constant(0.055)[0], due1=constant(0.055)[1]))):
        ws = wb.create_sheet(sheet)
        for row in rows:
            ws.append(row)
    wb.save(path)

    batch = load_workbook(db, path, people["editor"])
    assert batch.summary["agreements"] == 2 and batch.summary["proposed_rates"] == 2
    rates = db.scalars(select(RebateRate)).all()
    assert all(r.status is ReviewStatus.PROPOSED and r.entered_by_id == people["editor"].id for r in rates)
    costco = db.scalar(select(RebateAgreement).where(RebateAgreement.customer_label == "Costco"))
    assert costco.customer_group_code is None
    with pytest.raises(AlreadyLoaded):
        load_workbook(db, path, people["editor"])
    # A changed workbook for agreements already in the master: the platform is the master now.
    wb["JB Hi Fi"]["C1"] = "edited after the first load"
    changed = tmp_path / "MGNZ 2608 Aug 26 Rebates Reconciliation v2.xlsx"
    wb.save(changed)
    with pytest.raises(LoadRejected) as rejected:
        load_workbook(db, changed, people["editor"])
    assert "change rates in the platform" in rejected.value.problems[0]
