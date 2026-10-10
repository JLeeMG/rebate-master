"""Brand owners and their review packs: who checks which brand, the pack they get, and reading their answers."""

import io
from datetime import date
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import select

from mgrm.auth.roles import Role
from mgrm.data.brand_owners import load_brand_owners
from mgrm.data.errors import AlreadyLoaded, LoadRejected
from mgrm.models import AuditEvent, Brand, EvidenceFile
from mgrm.rebates.brand_review import ANSWERS, COLUMNS, HEADER_ROW, build_pack, load_return, owners, read_return
from tests.conftest import form_token, sign_in
from tests.test_accrual import load_june, pb, people  # noqa: F401  (fixtures)

JULY = date(2026, 7, 15)  # the twelve months to June 2026 include the PB Technology sales


def owners_file(tmp_path, rows, name="Brand owners.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Brand owners"
    ws.append(["Brand code", "Brand name", "Active", "Sales AU", "Sales NZ", "GPM name", "GPM email", "PM name", "PM email", "Notes"])
    for row in rows:
        ws.append(row)
    path = tmp_path / name
    wb.save(path)
    return path


GPMS = [["BONELK", "BONELK", "Yes", 0, 0, "David Vanderkley", "David@macgeargroup.com", "Kenneth Kan", "k.kan@macgeargroup.com", "Kenneth supports"],
        ["TWELVE_SOUTH", "TWELVE SOUTH", "Yes", 0, 0, "Rachel Harris", "r.harris@macgeargroup.com", "Prasanna Wad", "p.wad@macgeargroup.com", ""],
        ["SPACETALK", "SPACETALK", "Yes", 0, 0, "", "", "", "", "nobody"]]


@pytest.fixture
def owned(db, people, pb, tmp_path):
    load_june(db, people)
    load_brand_owners(db, owners_file(tmp_path, GPMS), people["siobhan"])
    return db


def test_owners_load_and_every_change_is_recorded(owned, db):
    bonelk = db.scalar(select(Brand).where(Brand.code == "BONELK"))
    assert (bonelk.gpm_name, bonelk.gpm_email, bonelk.pm_email) == ("David Vanderkley", "david@macgeargroup.com", "k.kan@macgeargroup.com")
    assert set(owners(db)) == {"david@macgeargroup.com", "r.harris@macgeargroup.com"}
    assert db.scalar(select(AuditEvent).where(AuditEvent.action == "brand.owner", AuditEvent.subject == "BONELK")) is not None


def test_packs_use_each_owners_actual_title(owned, db, people, tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Brand owners"
    ws.append(["Brand code", "GPM name", "GPM email", "GPM title", "PM name", "PM email", "Notes"])
    ws.append(["BONELK", "David Vanderkley", "david@macgeargroup.com", "Product Manager Lead - NZ", "", "", ""])
    path = tmp_path / "titles.xlsx"
    wb.save(path)
    load_brand_owners(db, path, people["siobhan"])
    _, content = build_pack(db, "david@macgeargroup.com", people["siobhan"], today=JULY)
    title = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Review"]["A1"].value
    assert title.endswith("David Vanderkley, Product Manager Lead - NZ")
    # A later file without the title column keeps the title on record
    later = [GPMS[0][:9] + ["Kenneth supports; reviewed Oct 2026"]] + GPMS[1:]
    load_brand_owners(db, owners_file(tmp_path, later, "again.xlsx"), people["siobhan"])
    assert db.scalar(select(Brand).where(Brand.code == "BONELK")).gpm_title == "Product Manager Lead - NZ"


@pytest.mark.parametrize(("row", "problem"), [
    (["BONELK", "", "", 0, 0, "Dave", "dave@gmail.com", "", "", ""], "not a MacGear email"),
    (["BONELK", "", "", 0, 0, "Dave", "", "", "", ""], "go together"),
    (["NOT_A_BRAND", "", "", 0, 0, "Dave", "d@macgeargroup.com", "", "", ""], "not in the brand register"),
    (["BONELK", "", "", 0, 0, "", "", "Kenneth", "k@macgeargroup.com", ""], "a PM without a GPM"),
])
def test_a_bad_owners_file_is_refused_whole(db, people, pb, tmp_path, row, problem):
    with pytest.raises(LoadRejected) as refused:
        load_brand_owners(db, owners_file(tmp_path, [row]), people["siobhan"])
    assert any(problem in p for p in refused.value.problems), refused.value.problems


def pack_rows(content: bytes) -> list[dict]:
    ws = openpyxl.load_workbook(io.BytesIO(content), read_only=True)["Review"]
    rows = list(ws.iter_rows(min_row=HEADER_ROW, values_only=True))
    return [dict(zip(COLUMNS, r)) for r in rows[1:] if r[0]]


def test_each_owner_gets_the_agreements_on_their_brands(owned, db, people, pb):
    _, content = build_pack(db, "david@macgeargroup.com", people["siobhan"], today=JULY)
    rows = pack_rows(content)
    assert [r["Ref"] for r in rows] == ["MGNZ-PBT-2"]  # Bonelk and Satechi at 10%; the 5% line excludes Bonelk
    assert rows[0]["Your brands"] == "BONELK" and rows[0]["Rate"] == "10%" and rows[0]["Type"] == "Rebate"
    assert rows[0]["What the agreement covers"] == "Only BONELK, SATECHI"  # no journal tagging, no internal ids
    assert rows[0]["Your brands' sales, last 12 months"] == pytest.approx(61390.95)
    # Twelve South is covered by the retailer-wide 5% line ("all brands except ...")
    _, content = build_pack(db, "r.harris@macgeargroup.com", people["siobhan"], today=JULY)
    assert [r["Ref"] for r in pack_rows(content)] == ["MGNZ-PBT-1"]
    assert db.scalar(select(AuditEvent).where(AuditEvent.action == "review_pack.issue")) is not None


def test_the_pack_shows_owners_only_what_they_need(owned, db, people):
    _, content = build_pack(db, "david@macgeargroup.com", people["siobhan"], today=JULY)
    wb = openpyxl.load_workbook(io.BytesIO(content))
    review = wb["Review"]
    assert review.column_dimensions["A"].hidden  # the agreement reference
    assert wb["About this pack"].sheet_state == "veryHidden"
    text = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value)
    for internal in ("journal", "Kenneth supports", "Source", "Accrued", "awaiting approval", "approved by Finance",
                     "platform", "Rebate Master"):
        assert internal not in text, internal  # Finance's notes and working columns stay with Finance


def answered(content: bytes, tmp_path, answers: dict[str, tuple], name="returned.xlsx"):
    """The owner fills in the yellow columns and saves (as Excel would)."""
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb["Review"]
    for row in range(HEADER_ROW + 1, ws.max_row + 1):
        ref = ws.cell(row=row, column=1).value
        for offset, value in enumerate(answers.get(ref, ())):
            ws.cell(row=row, column=COLUMNS.index(ANSWERS[0]) + 1 + offset, value=value)
    path = tmp_path / name
    wb.save(path)
    return path


def test_a_returned_pack_records_confirmations_and_lists_corrections(owned, db, people, pb, tmp_path):
    _, content = build_pack(db, "david@macgeargroup.com", people["siobhan"], today=JULY)
    path = answered(content, tmp_path, {"MGNZ-PBT-2": ("N", 12, "New terms from 1 Jan, email attached to Siobhan")})
    batch, result = load_return(db, path, people["siobhan"])
    assert result.owner == "david@macgeargroup.com"
    assert result.corrections == [("MGNZ-PBT-2", Decimal("0.120000"), "New terms from 1 Jan, email attached to Siobhan")]
    assert any("MGNZ-PBT-2: says 12%" in n for n in batch.summary["notes"])
    evidence = db.scalar(select(EvidenceFile).where(EvidenceFile.agreement_id == pb["MGNZ-PBT-2"].id))
    assert evidence.file_name == "returned.xlsx" and evidence.content_type.endswith("spreadsheetml.sheet")
    assert pb["MGNZ-PBT-2"].scope is not None and db.get(type(pb["MGNZ-PBT-2"]), pb["MGNZ-PBT-2"].id)  # nothing changed
    with pytest.raises(AlreadyLoaded):
        load_return(db, path, people["siobhan"])


def test_a_confirmation_is_recorded_against_the_agreement(owned, db, people, pb, tmp_path):
    _, content = build_pack(db, "r.harris@macgeargroup.com", people["siobhan"], today=JULY)
    _, result = load_return(db, answered(content, tmp_path, {"MGNZ-PBT-1": ("Y",)}), people["siobhan"])
    assert result.confirmed == ["MGNZ-PBT-1"]
    event = db.scalar(select(AuditEvent).where(AuditEvent.action == "rebate.owner_confirmed"))
    assert event.subject == "MGNZ-PBT-1" and event.detail["owner"] == "r.harris@macgeargroup.com"


def test_a_pack_with_rows_removed_is_refused(owned, db, people, pb, tmp_path):
    _, content = build_pack(db, "david@macgeargroup.com", people["siobhan"], today=JULY)
    wb = openpyxl.load_workbook(io.BytesIO(content))
    wb["Review"].delete_rows(HEADER_ROW + 1)
    path = tmp_path / "edited.xlsx"
    wb.save(path)
    _, problems = read_return(path)
    assert any("re-issued" in p for p in problems)


def test_the_brands_page_offers_packs_to_finance_only(client, owned, people, make_user):
    sign_in(client, "siobhan@example.com")
    page = client.get("/brands").text
    assert "David Vanderkley" in page and "Download review pack" in page
    download = client.get("/brands/review-pack?owner=david%40macgeargroup.com")
    assert download.status_code == 200 and download.headers["content-type"].startswith("application/vnd.openxmlformats")
    client.post("/logout", data={"csrf_token": form_token(client, "/")})
    sign_in(client, "raymond@example.com")
    assert "Download review pack" not in client.get("/brands").text
    assert client.get("/brands/review-pack?owner=david%40macgeargroup.com").status_code == 403
