"""The screens: who sees what, the evidence flow end to end, and L1/L3 on every page."""

import openpyxl
import pytest

from mgrm.auth.roles import Role
from mgrm.checks.text import scan_generated_text
from mgrm.data.files import bytes_sha256
from mgrm.models import EvidenceFile, RebateAgreement, RebateRate, ReviewStatus
from tests.conftest import PDF_BYTES, form_token, sign_in, switch
from tests.test_rebates import constant, tab_rows

READ_PAGES = ["/", "/rebates", "/rebates/change-log", "/rebates/expiry", "/rebates/rate-card", "/customers", "/brands"]
REVIEW, ADMIN_PAGES, LOADS = ["/rebates/review"], ["/admin/users", "/admin/audit", "/admin/feed-tokens", "/rebates/approvers"], ["/loads"]


@pytest.mark.parametrize(
    ("role", "allowed", "refused"),
    [
        (Role.VIEWER, READ_PAGES, REVIEW + ADMIN_PAGES + LOADS),
        (Role.BRAND_APPROVER, READ_PAGES + REVIEW, ADMIN_PAGES + LOADS),
        (Role.REBATE_REVIEWER, READ_PAGES + REVIEW, ADMIN_PAGES + LOADS),
        (Role.REBATE_MAINTAINER, READ_PAGES + LOADS + REVIEW + ["/rebates/new"], ADMIN_PAGES),
        (Role.REBATE_EDITOR, READ_PAGES + LOADS + ["/rebates/new"], REVIEW + ADMIN_PAGES),
        (Role.ADMIN, READ_PAGES + REVIEW + ADMIN_PAGES + LOADS, []),
    ],
)
def test_page_access(client, make_user, role, allowed, refused):
    make_user(role)
    sign_in(client, f"{role.value}@example.com")
    for path in allowed:
        assert client.get(path).status_code == 200, path
    for path in refused:
        assert client.get(path).status_code == 403, path


def test_pages_need_sign_in(client):
    for path in READ_PAGES[1:] + ["/rebates/evidence/1"]:
        response = client.get(path, follow_redirects=False)
        assert response.status_code == 303 and response.headers["location"] == "/login", path


def workbook(tmp_path) -> bytes:
    path = tmp_path / "MGNZ 2608 Aug 26 Rebates Reconciliation.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Financial Summary"
    ws = wb.create_sheet("NLG 20%")
    sales, due = constant(0.2)
    for row in tab_rows(customer="Noel Leeming", rate1=0.2, sales1=sales, due1=due):
        ws.append(row)
    wb.save(path)
    return path.read_bytes()


def load_workbook_through_screen(client, tmp_path):
    return client.post("/loads", data={"kind": "rebate_workbook", "csrf_token": form_token(client, "/loads")},
                       files={"upload": ("MGNZ 2608 Aug 26 Rebates Reconciliation.xlsx", workbook(tmp_path))})


def test_siobhan_proposes_with_evidence_ken_reviews_raymond_reads(client, make_user, db, tmp_path):
    make_user(Role.REBATE_EDITOR, email="siobhan@example.com", display_name="Siobhan")
    make_user(Role.REBATE_REVIEWER, email="ken@example.com", display_name="Ken")
    make_user(Role.VIEWER, email="raymond@example.com", display_name="Raymond")

    sign_in(client, "siobhan@example.com")
    assert "Loaded MGNZ 2608" in load_workbook_through_screen(client, tmp_path).text
    agreement = db.query(RebateAgreement).one()
    page = f"/rebates/agreement/{agreement.id}"

    without = client.post(f"{page}/rates", data={
        "csrf_token": form_token(client, page), "rate_percent": "25.3", "effective_from": "2026-10-01",
        "source_reference": "Email from NLG", "reason": "FY27 terms"})
    assert without.status_code == 400 and "Attach the evidence" in without.text

    proposed = client.post(f"{page}/rates", data={
        "csrf_token": form_token(client, page), "rate_percent": "25.3", "effective_from": "2026-10-01",
        "source_reference": "Email from NLG, 2 Oct 2026", "reason": "FY27 terms"},
        files=[("evidence", ("NLG agreement email.pdf", PDF_BYTES, "application/pdf"))])
    assert "proposed with its evidence" in proposed.text
    new_rate = db.query(RebateRate).filter(RebateRate.reason == "FY27 terms").one()
    evidence = db.query(EvidenceFile).one()

    switch(client, "ken@example.com")
    queue = client.get("/rebates/review").text
    assert "FY27 terms" in queue and "NLG agreement email.pdf" in queue
    opened = client.get(f"/rebates/evidence/{evidence.id}")
    assert opened.content == PDF_BYTES and opened.headers["x-content-sha256"] == bytes_sha256(PDF_BYTES)
    client.post(f"/rebates/rates/{new_rate.id}/decision",
                data={"csrf_token": form_token(client, "/rebates/review"), "decision": "approve", "note": "matches the email"})
    db.refresh(new_rate)
    assert new_rate.status is ReviewStatus.APPROVED

    switch(client, "raymond@example.com")
    log = client.get("/rebates/change-log").text
    for expected in ("Siobhan", "Ken", "FY27 terms", "25.3%", "NLG agreement email.pdf", "matches the email", "Approved"):
        assert expected in log, expected
    detail = client.get(page).text
    assert "Rate applied" in detail and "NZ$" not in detail  # a viewer sees rates, not sales amounts
    assert "Propose a rate change" not in detail


def test_feed_token_is_shown_once(client, make_user, db):
    make_user(Role.ADMIN)
    sign_in(client, "admin@example.com")
    created = client.post("/admin/feed-tokens", data={"csrf_token": form_token(client, "/admin/feed-tokens"), "name": "Forecasting platform"})
    assert "will not be shown again" in created.text
    assert "will not be shown again" not in client.get("/admin/feed-tokens").text


def test_every_page_is_free_of_bare_dollars_and_markers(client, make_user, db, tmp_path):
    make_user(Role.ADMIN)
    sign_in(client, "admin@example.com")
    load_workbook_through_screen(client, tmp_path)
    agreement = db.query(RebateAgreement).one()
    for path in READ_PAGES + REVIEW + ADMIN_PAGES + LOADS + [f"/rebates/agreement/{agreement.id}"]:
        response = client.get(path)
        assert response.status_code == 200, path
        assert scan_generated_text(response.text) == [], path


def test_ken_and_siobhan_enter_and_approve_each_others_input(client, make_user, db):
    from mgrm.models import RebateChangeRequest

    make_user(Role.REBATE_MAINTAINER, email="siobhan@example.com", display_name="Siobhan")
    make_user(Role.REBATE_MAINTAINER, email="ken@example.com", display_name="Ken")

    sign_in(client, "siobhan@example.com")
    created = client.post("/rebates/new", data={
        "csrf_token": form_token(client, "/rebates/new"), "entity": "MGNZ", "customer_label": "PB Technology",
        "customer_group_code": "PBT", "brand_code": "EUFY", "product_scope": "EUFY", "rate_type": "rebate",
        "basis": "rebate_eligible_sales", "accrual_mode": "accrue", "rate_percent": "15", "effective_from": "2026-11-01",
        "source_reference": "Email from PB Tech", "reason": "EUFY ranged at PB Tech"},
        files=[("evidence", ("PB email.pdf", PDF_BYTES, "application/pdf"))])
    assert "New agreement proposed" in created.text
    agreement = db.query(RebateAgreement).one()
    queue = client.get("/rebates/review").text
    assert "new agreement" in queue and "someone else must review it" in queue  # her own: no approve button

    switch(client, "ken@example.com")
    rate = db.query(RebateRate).one()
    client.post(f"/rebates/rates/{rate.id}/decision", data={"csrf_token": form_token(client, "/rebates/review"), "decision": "approve"})
    db.refresh(rate)
    assert rate.status is ReviewStatus.APPROVED

    page = f"/rebates/agreement/{agreement.id}"
    proposed = client.post(f"{page}/end", data={"csrf_token": form_token(client, page), "effective_to": "2027-03-31",
                                                "reason": "Range delisted"})
    assert "Ending proposed" in proposed.text
    change = db.query(RebateChangeRequest).one()

    switch(client, "siobhan@example.com")
    client.post(f"/rebates/changes/{change.id}/decision", data={"csrf_token": form_token(client, "/rebates/review"), "decision": "approve"})
    db.refresh(rate)
    assert rate.effective_to.isoformat() == "2027-03-31"
    log = client.get("/rebates/change-log").text
    for expected in ("Proposed a new agreement", "Proposed ending the agreement", "Approved the ending", "Range delisted"):
        assert expected in log, expected
