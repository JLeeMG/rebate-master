"""One-time load of the rebate master from the legacy reconciliation workbooks (spec §4.5.1, §5.3).

Both entities' workbooks ("MGAU 2608 Aug 26 Rebates Reconciliation - Actual WP
Final.xlsx", "MGNZ 2608 Aug 26 Rebates Reconciliation.xlsx") hold one tab per
agreement on a shared template:

    rows 1-4    customer, product groups, rates, and free-text notes
    month row   one column per month since 2017 (MGAU) or 2021 (MGNZ)
    Sales Ex GST    (two lines)     sales per product group, per month
    Rebate Due      (two lines)     rebate accrued per product group; column B holds the rate
    Total Sales / Total Rebate Due / credits / over-(under) accrual

Each tab's rate is a single cell that was overwritten when it changed, so the
history is rebuilt from what was actually applied: rebate due / sales, month
by month. A period is proposed as a rate only where that applied rate is also
written somewhere on the tab ("confirmed"). Anything else - blended rates,
impossible rates, months with sales but no rebate - is carried as a flag for
the reviewer, never guessed.

Everything loads as PROPOSED. Nothing is used until a second person approves it.
Workbooks are opened read-only and never saved (spec §14).
"""

import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from mgrm.data.errors import AlreadyLoaded, LoadRejected
from mgrm.data.files import file_sha256
from mgrm.domain.entities import Entity
from mgrm.domain.periods import add_months
from mgrm.models import (
    AccrualMode,
    AppUser,
    LoadBatch,
    LoadKind,
    RateType,
    RebateAgreement,
    RebateBasis,
    RebateRate,
    RebateWorkbookMonth,
    ReviewStatus,
)

HEADER_ROWS = 30
MIN_MONTH_COLUMNS = 6
LINES_PER_TAB = 2
# Two rates are the same rate if they agree to a thousandth of a percentage point.
RATE_TOLERANCE = Decimal("0.00001")
# An applied rate outside this range is not a rate; it is a data error in the tab.
MAX_PLAUSIBLE_RATE = Decimal("0.6")
# Sales below this are noise (rounding, a stray credit) and say nothing about the rate.
MIN_SALES_FOR_RATE = Decimal("1")
# No sales for this many months before the workbook's month: treated as ended, and flagged.
MONTHS_WITHOUT_SALES_TO_END = 3
MONEY = Decimal("0.01")
# A month applied a written rate if rate x sales is within this of the rebate due (cent rounding,
# sometimes across two lines). Otherwise the applied rate is recorded as unconfirmed.
ROUNDING_ALLOWANCE = Decimal("0.05")

NOT_AGREEMENTS = {
    "customer summary", "financial summary", "start", "end", "master rebate sheet", "additional",
    "rebates accruals", "rebates check vs bp 03-26", "jbhifi check 2% & algo",
}
CHECK_ONLY_MARKER = "purely for check"  # the Honor tabs: "DO NOT ACCRUE PURELY FOR CHECK"
EXCLUSION_MARKER = "do not accrue"  # anything else saying this excludes part of the sales; flagged, still accrued
# Customer name fragments -> customer group code. Anything unmatched loads with no group and is flagged.
GROUP_KEYWORDS = (
    ("harvey norman", "HN"), ("hn ", "HN"), ("jb hi", "JBH"), ("jbhifi", "JBH"), ("amazon", "AMZ"),
    ("nlg", "NL"), ("noel leeming", "NL"), ("officeworks", "OW"), ("good guys", "TGG"), ("pbt", "PBT"),
    ("pb tech", "PBT"),
)
RATE_TYPE_KEYWORDS = (("mdf", RateType.MDF), ("dam", RateType.DAMAGE), ("sett", RateType.SETTLEMENT), ("co-op", RateType.CO_OP))
START_HINT = re.compile(r"(?:(\d{1,2})/)?(\d{1,2})/(\d{4}|\d{2})\s*onwards", re.IGNORECASE)
FILE_PERIOD = re.compile(r"\b(\d{2})(\d{2}) [A-Z][a-z]{2} \d{2}\b")
ENTITY_IN_NAME = re.compile(r"^(MGAU|MGNZ)\b")


@dataclass(frozen=True)
class Run:
    start: date
    end: date  # first day of the last month in the run
    rate: Decimal
    confirmed: bool


@dataclass
class Line:
    number: int  # 1 or 2, the template's product-group line
    scope: str
    rate_cell: Decimal | None
    months: dict[date, tuple[Decimal, Decimal]] = field(default_factory=dict)  # period -> (sales, rebate due)
    runs: list[Run] = field(default_factory=list)
    flags: list[str] = field(default_factory=list)


@dataclass
class Tab:
    entity: Entity
    sheet: str
    customer_label: str
    group_code: str | None
    rate_type: RateType
    accrual_mode: AccrualMode
    notes: list[str]
    start_hint: date | None
    lines: list[Line]
    flags: list[str] = field(default_factory=list)


@dataclass
class ParsedWorkbook:
    entity: Entity | None
    as_at: date | None
    tabs: list[Tab]
    skipped: list[str]
    problems: list[str]


RATE_PLACES = Decimal("0.000001")  # the database holds rates to six decimal places


def _decimal(value) -> Decimal | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return Decimal(repr(value))
    return None


def _rate(value) -> Decimal | None:
    """A rate as written in Excel, without binary noise (0.14500000000000002 -> 0.145)."""
    number = _decimal(value)
    return None if number is None else number.quantize(RATE_PLACES).normalize()


def _text(value) -> str:
    if isinstance(value, str):
        return value.strip()
    return ""


def _label(row: list) -> str:
    return _text(row[0]) if row else ""


def _pct(rate: Decimal) -> str:
    return f"{(rate * 100).normalize():f}%"


def _start_hint(notes: list[str]) -> date | None:
    for note in notes:
        match = START_HINT.search(note)
        if match:
            _day, month, year = match.groups()
            year_value = int(year) if len(year) == 4 else 2000 + int(year)
            if 1 <= int(month) <= 12:
                return date(year_value, int(month), 1)
    return None


def _group_for(*names: str) -> str | None:
    text = " ".join(names).lower() + " "
    for fragment, code in GROUP_KEYWORDS:
        if fragment in text:
            return code
    return None


def _rate_type(*names: str) -> RateType:
    text = " ".join(names).lower()
    for fragment, rate_type in RATE_TYPE_KEYWORDS:
        if fragment in text:
            return rate_type
    return RateType.REBATE


def _same(a: Decimal, b: Decimal) -> bool:
    return abs(a - b) <= RATE_TOLERANCE


def build_runs(line: Line, written_rates: set[Decimal]) -> None:
    """Group successive months applying the same rate. Months without sales do not break a run."""
    anomalies, no_rebate = [], []
    applied: list[tuple[date, Decimal, bool]] = []
    for period in sorted(line.months):
        sales, due = line.months[period]
        if abs(sales) < MIN_SALES_FOR_RATE:
            continue
        if due == 0:
            no_rebate.append(period)
            continue
        written = next((w for w in sorted(written_rates) if abs(due - w * sales) <= ROUNDING_ALLOWANCE), None)
        if written is not None:
            applied.append((period, written, True))
            continue
        rate = (due / sales).quantize(RATE_PLACES)
        if rate < 0 or rate >= MAX_PLAUSIBLE_RATE:
            anomalies.append(f"{period:%b %Y} ({_pct(rate)})")
            continue
        applied.append((period, rate, False))
    runs: list[Run] = []
    for period, rate, confirmed in applied:
        if runs and _same(runs[-1].rate, rate) and runs[-1].confirmed == confirmed:
            runs[-1] = Run(runs[-1].start, period, runs[-1].rate, confirmed)
        else:
            runs.append(Run(period, period, rate, confirmed))
    line.runs = runs
    if anomalies:
        line.flags.append(f"Impossible applied rate, excluded: {', '.join(anomalies)}")
    if no_rebate:
        line.flags.append(
            f"{len(no_rebate)} month(s) with sales but no rebate due, e.g. {', '.join(f'{p:%b %Y}' for p in no_rebate[:4])}"
        )
    unconfirmed = [r for r in runs if not r.confirmed]
    if unconfirmed:
        sample = "; ".join(f"{r.start:%b %Y}-{r.end:%b %Y} {_pct(r.rate)}" for r in unconfirmed[:4])
        more = f" and {len(unconfirmed) - 4} more" if len(unconfirmed) > 4 else ""
        line.flags.append(
            f"{len(unconfirmed)} period(s) applied a rate not written on the tab (blended or manual): {sample}{more}"
        )


def parse_tab(rows: list[list], sheet: str, entity: Entity) -> Tab | str:
    """A Tab, or a string saying why the sheet was skipped."""
    month_row = next((i for i, r in enumerate(rows) if sum(isinstance(c, datetime) for c in r) > MIN_MONTH_COLUMNS), None)
    if month_row is None:
        return "no month columns: not an agreement tab"
    months = {j: c.date().replace(day=1) for j, c in enumerate(rows[month_row]) if isinstance(c, datetime)}
    sales_rows = [i for i, r in enumerate(rows) if _label(r).startswith("Sales Ex GST")][:LINES_PER_TAB]
    due_rows = [i for i, r in enumerate(rows) if _label(r).startswith("Rebate Due")][:LINES_PER_TAB]
    if len(sales_rows) != LINES_PER_TAB or len(due_rows) != LINES_PER_TAB:
        return "the Sales Ex GST / Rebate Due rows are not where the template puts them"

    top = rows[:month_row]
    notes = []
    for row in top:
        for cell in row:
            text = _text(cell)
            if text and text not in notes and text not in ("Home", "Product Groups", "MacGear"):
                notes.append(text)
    written_rates = {d for row in top for c in row if (d := _rate(c)) is not None and 0 < d < MAX_PLAUSIBLE_RATE}
    customer_label = _text(rows[2][0]) if len(rows) > 2 and rows[2] and _text(rows[2][0]) else sheet.strip()

    lines = []
    for number, (sales_row, due_row) in enumerate(zip(sales_rows, due_rows), start=1):
        rate_cell = _rate(rows[due_row][1]) if len(rows[due_row]) > 1 else None
        if rate_cell is not None and rate_cell <= 0:
            rate_cell = None
        if rate_cell is not None:
            written_rates.add(rate_cell)
        scope = _text(rows[sales_row][1]) if len(rows[sales_row]) > 1 else ""
        for fallback in ((2, 2), (3, 2), (2, 5), (3, 5)) if not scope else ():
            r, c = fallback
            if r == number + 1 and len(rows) > r and len(rows[r]) > c and _text(rows[r][c]):
                scope = _text(rows[r][c])
                break
        line = Line(number, scope or ("ALL (not stated on the tab)" if number == 1 else "second line (not stated)"), rate_cell)
        for j, period in months.items():
            sales = _decimal(rows[sales_row][j]) if j < len(rows[sales_row]) else None
            due = _decimal(rows[due_row][j]) if j < len(rows[due_row]) else None
            if (sales or 0) or (due or 0):
                line.months[period] = ((sales or Decimal(0)).quantize(MONEY), (due or Decimal(0)).quantize(MONEY))
        if rate_cell is None and not line.months:
            continue  # an unused second line
        lines.append(line)
    if not lines:
        return "blank: no rates and no figures"
    for line in lines:
        build_runs(line, written_rates)

    tab = Tab(
        entity=entity,
        sheet=sheet.strip(),
        customer_label=customer_label,
        group_code=_group_for(customer_label, sheet),
        rate_type=_rate_type(sheet, *[l.scope for l in lines]),
        accrual_mode=AccrualMode.CHECK_ONLY if any(CHECK_ONLY_MARKER in n.lower() for n in notes) else AccrualMode.ACCRUE,
        notes=notes,
        start_hint=_start_hint(notes),
        lines=lines,
    )
    if tab.accrual_mode is AccrualMode.CHECK_ONLY:
        marker = next(n for n in notes if CHECK_ONLY_MARKER in n.lower())
        tab.flags.append(f"Loaded as check-only (not accrued) because the tab says: '{marker}'")
    else:
        for note in notes:
            if EXCLUSION_MARKER in note.lower():
                tab.flags.append(f"The tab says '{note}': check whether part of the sales must be excluded from the base")
    if tab.group_code is None:
        tab.flags.append(f"No customer group for '{customer_label}': choose or create one before approving")
    return tab


def parse_workbook(path: Path) -> ParsedWorkbook:
    problems: list[str] = []
    match = ENTITY_IN_NAME.match(path.name)
    entity = Entity(match.group(1)) if match else None
    if entity is None:
        problems.append("the file name must start with MGAU or MGNZ")
    period = FILE_PERIOD.search(path.name)
    as_at = date(2000 + int(period.group(1)), int(period.group(2)), 1) if period else None
    if as_at is None:
        problems.append("the file name must carry its month as YYMM, e.g. '2608 Aug 26'")
    if problems:
        return ParsedWorkbook(entity, as_at, [], [], problems)

    tabs, skipped = [], []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            if ws.title.strip().lower() in NOT_AGREEMENTS:
                skipped.append(f"{ws.title.strip()}: summary or template sheet")
                continue
            rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=HEADER_ROWS, values_only=True)]
            parsed = parse_tab(rows, ws.title, entity)
            if isinstance(parsed, str):
                skipped.append(f"{ws.title.strip()}: {parsed}")
            else:
                tabs.append(parsed)
    finally:
        wb.close()
    if not tabs:
        problems.append("no agreement tabs found")
    return ParsedWorkbook(entity, as_at, tabs, skipped, problems)


# ---------------------------------------------------------------- proposals


@dataclass(frozen=True)
class ProposedRate:
    rate: Decimal
    effective_from: date
    effective_to: date | None
    flags: tuple[str, ...]


def _month_end(period: date) -> date:
    return add_months(period, 1) - timedelta(days=1)


def propose_rates(tab: Tab, line: Line, as_at: date) -> list[ProposedRate]:
    """Confirmed runs become rate periods, each running until the next begins.

    The rate applied in the latest months is the current rate. A rate cell that
    disagrees with it is flagged, not proposed: on OfficeWorks the column-B cell
    still reads 10% while 19% has been applied since July 2026.
    """
    confirmed: list[Run] = []
    for run in (r for r in line.runs if r.confirmed):
        if confirmed and _same(confirmed[-1].rate, run.rate):
            confirmed[-1] = Run(confirmed[-1].start, run.end, run.rate, True)  # rejoin across a one-off month
        else:
            confirmed.append(run)

    active_until = max(line.months) if line.months else None
    ended = active_until is not None and active_until <= add_months(as_at, -MONTHS_WITHOUT_SALES_TO_END)
    proposals: list[ProposedRate] = []
    for index, run in enumerate(confirmed):
        following = confirmed[index + 1] if index + 1 < len(confirmed) else None
        end = _month_end(add_months(following.start, -1)) if following else None
        proposals.append(ProposedRate(run.rate, run.start, end, ()))

    current = line.rate_cell
    if proposals:
        last = proposals[-1]
        if current is not None and not _same(last.rate, current):
            proposals[-1] = ProposedRate(last.rate, last.effective_from, last.effective_to, last.flags + (
                f"The rate cell on the tab says {_pct(current)} but {_pct(last.rate)} was applied in the latest "
                f"months; {_pct(last.rate)} is proposed",))
    elif current is not None:
        hint = _start_hint([line.scope]) or tab.start_hint
        start = hint or (min(line.months) if line.months else as_at)
        source = "the tab's note" if hint else ("the first month with figures" if line.months else "the workbook's month")
        proposals.append(ProposedRate(current, start, None, (
            f"No month confirms this rate; it is the rate written on the tab, from {source}",)))
    else:
        line.flags.append("No rate could be established from this line")

    if ended and proposals:
        last = proposals[-1]
        note = f"No sales since {active_until:%b %Y}: proposed as ended then"
        proposals[-1] = ProposedRate(last.rate, last.effective_from, _month_end(active_until), last.flags + (note,))
    return proposals


def agreement_code(entity: Entity, sheet: str, line: int) -> str:
    slug = re.sub(r"[^A-Z0-9]+", "-", sheet.upper()).strip("-")
    return f"{entity.value}-{slug}-{line}"


def load_workbook(session: Session, path: Path, actor: AppUser) -> LoadBatch:
    """Save every tab's agreements and proposed rates. Raises LoadRejected or AlreadyLoaded."""
    path = Path(path)
    parsed = parse_workbook(path)
    if parsed.problems:
        raise LoadRejected(path.name, parsed.problems)
    sha = file_sha256(path)
    if session.scalar(select(LoadBatch).where(LoadBatch.kind == LoadKind.REBATE_WORKBOOK, LoadBatch.file_sha256 == sha)):
        raise AlreadyLoaded(f"{path.name} has already been loaded")
    codes = [agreement_code(parsed.entity, t.sheet, l.number) for t in parsed.tabs for l in t.lines]
    existing = set(session.scalars(select(RebateAgreement.code).where(RebateAgreement.code.in_(codes))))
    if existing:
        raise LoadRejected(path.name, [
            f"{len(existing)} of its agreements are already in the rebate master (e.g. {sorted(existing)[0]}). "
            "After the first load the platform is the master: change rates in the platform, not by reloading."
        ])

    batch = LoadBatch(
        kind=LoadKind.REBATE_WORKBOOK, entity_code=parsed.entity.value, source="legacy rebate reconciliation workbook",
        file_name=path.name, file_sha256=sha, row_count=0, loaded_by_id=actor.id,
        summary={"as_at": parsed.as_at.isoformat(), "skipped": parsed.skipped},
    )
    session.add(batch)
    session.flush()
    agreements = rates = flagged = 0
    for tab in parsed.tabs:
        for line in tab.lines:
            source = f"{path.name} / tab '{tab.sheet}' / line {line.number}"
            agreement = RebateAgreement(
                code=agreement_code(parsed.entity, tab.sheet, line.number),
                entity_code=parsed.entity.value,
                customer_label=tab.customer_label,
                customer_group_code=tab.group_code,
                brand_code="ALL" if line.scope.upper().startswith("ALL") else None,
                product_scope=line.scope,
                rate_type=tab.rate_type,
                basis=RebateBasis.REBATE_ELIGIBLE_SALES,
                accrual_mode=tab.accrual_mode,
                source_reference=source,
                notes="\n".join(tab.notes),
                created_by_id=actor.id,
            )
            session.add(agreement)
            session.flush()
            agreements += 1
            tab_flags = tab.flags + line.flags
            for proposal in propose_rates(tab, line, parsed.as_at):
                flags = list(tab_flags) + list(proposal.flags)
                flagged += bool(flags)
                session.add(RebateRate(
                    agreement_id=agreement.id, rate=proposal.rate, effective_from=proposal.effective_from,
                    effective_to=proposal.effective_to, status=ReviewStatus.PROPOSED, source_reference=source,
                    reason="Initial load from the legacy reconciliation workbook", load_flags=flags,
                    entered_by_id=actor.id,
                ))
                rates += 1
            session.add_all(
                RebateWorkbookMonth(batch_id=batch.id, agreement_id=agreement.id, period=p, sales=s, rebate_due=d)
                for p, (s, d) in line.months.items()
            )
    batch.row_count = rates
    batch.summary = {**batch.summary, "agreements": agreements, "proposed_rates": rates, "flagged_rates": flagged,
                     "tabs": len(parsed.tabs)}
    session.flush()
    return batch
