"""Brand review packs: each brand owner (GPM) checks the rebate agreements on their brands, in Excel.

While the platform is used by Finance only, owners don't sign in. Finance sends each GPM a workbook
listing every agreement that touches their brands (including retailer-wide agreements that cover
them), with the rate in force, any change awaiting approval, and their brands' sales at that
retailer. They answer three columns: Correct? (Y/N), Correct rate (%), Comment.

When the workbook comes back it is read (never edited or saved): confirmations are recorded against
each agreement in the audit log; corrections and comments are listed for Finance, who enter any
change as an ordinary proposal for someone else to approve, with the returned workbook as evidence.
"""

import io
import json
from collections import defaultdict
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from mgrm.auth.users import audit
from mgrm.data.errors import AlreadyLoaded, LoadRejected
from mgrm.data.files import file_sha256, workbook_problem
from mgrm.models import (
    AccrualMode,
    AppUser,
    Brand,
    Customer,
    LoadBatch,
    LoadKind,
    RebateAgreement,
    RebateChangeRequest,
    RebateRate,
    ReviewStatus,
    SalesLine,
)
from mgrm.rebates.accrual import month_end
from mgrm.rebates.scope import ScopeError, parse_scope
from mgrm.rebates.service import Upload, attach_evidence, rate_in_force

SHEET = "Review"
ABOUT = "About this pack"
HEADER_ROW = 6
COLUMNS = ["Ref", "Entity", "Retailer", "Your brands it covers", "What the agreement covers", "Type", "Accrued?",
           "Rate in force", "Since", "Until", "Change awaiting approval", "Source", "Your brands' sales, last 12 months",
           "Correct? (Y/N)", "Correct rate (%)", "Comment"]
ANSWERS = ("Correct? (Y/N)", "Correct rate (%)", "Comment")
SALES_MONTHS = 12
CURRENCY = {"MGAU": "A$", "MGNZ": "NZ$"}


@dataclass(frozen=True)
class PackLine:
    agreement: RebateAgreement
    brands: list[str]
    scope_text: str
    scope_confirmed: bool
    rate: RebateRate | None
    pending: RebateRate | None
    sales: Decimal


def owners(session: Session) -> dict[str, dict]:
    """GPM email -> {name, brands, pm emails}."""
    out: dict[str, dict] = {}
    for b in session.scalars(select(Brand).where(Brand.gpm_email != "").order_by(Brand.code)):
        entry = out.setdefault(b.gpm_email, {"name": b.gpm_name, "brands": [], "cc": set(), "notes": {}})
        entry["brands"].append(b.code)
        if b.pm_email and b.pm_email != b.gpm_email:
            entry["cc"].add(b.pm_email)
        if b.owner_notes:
            entry["notes"][b.code] = b.owner_notes
    return out


def _last_closed_month(today: date) -> date:
    first = today.replace(day=1)
    return date(first.year - (first.month == 1), first.month - 1 or 12, 1)


def pack_lines(session: Session, brands: list[str], today: date | None = None) -> list[PackLine]:
    today = today or date.today()
    last = _last_closed_month(today)
    first = date(last.year - 1 + (last.month == 12), last.month % 12 + 1, 1)
    code_of = {b.netsuite_class_id: b.code for b in session.scalars(select(Brand)) if b.netsuite_class_id is not None}
    group_of = {c.netsuite_customer_id: c.customer_group_code for c in session.scalars(select(Customer))}
    sales: dict[tuple[str, str, str], Decimal] = defaultdict(Decimal)  # (entity, group, brand) -> 12 months' sales
    for class_id, customer_id, entity, amount in session.execute(
            select(SalesLine.netsuite_class_id, SalesLine.netsuite_customer_id, SalesLine.entity_code, func.sum(SalesLine.amount))
            .where(SalesLine.period >= first, SalesLine.period <= last)
            .group_by(SalesLine.netsuite_class_id, SalesLine.netsuite_customer_id, SalesLine.entity_code)):
        group, code = group_of.get(customer_id), code_of.get(class_id)
        if group and code in brands:
            sales[(entity, group, code)] += amount
    pending_scope = {}
    for change in session.scalars(select(RebateChangeRequest).where(RebateChangeRequest.status == ReviewStatus.PROPOSED)
                                  .order_by(RebateChangeRequest.id)):
        new = change.payload.get("changes", {}).get("scope")
        if new:
            pending_scope[change.agreement_id] = new[1]
    rates = defaultdict(list)
    for r in session.scalars(select(RebateRate)):
        rates[r.agreement_id].append(r)

    lines = []
    for a in session.scalars(select(RebateAgreement).order_by(RebateAgreement.entity_code, RebateAgreement.customer_label,
                                                              RebateAgreement.code)):
        text = a.scope or pending_scope.get(a.id)
        try:
            scope = parse_scope(text)
        except ScopeError:
            scope = None
        if scope is None:
            covered = [b for b in brands if a.brand_code == b]
        elif scope.brand_mode == "only":
            covered = sorted(set(brands) & scope.brand_codes)
        else:  # retailer-wide: list only the owner's brands this retailer actually buys
            covered = [b for b in brands if scope.includes_brand(b) and sales.get((a.entity_code, a.customer_group_code, b))]
        if not covered:
            continue
        in_force = rate_in_force(rates[a.id], today)
        pending = max((r for r in rates[a.id] if r.status is ReviewStatus.PROPOSED), key=lambda r: r.id, default=None)
        if in_force is None and pending is None:
            continue  # ended, or never in force
        lines.append(PackLine(a, covered, scope.describe() if scope else a.product_scope, a.scope is not None,
                              in_force, pending,
                              sum((sales.get((a.entity_code, a.customer_group_code, b), Decimal(0)) for b in covered), Decimal(0))))
    return lines


def _pct(rate: Decimal | None) -> str:
    return "" if rate is None else f"{(rate * 100).normalize():f}%"


def build_pack(session: Session, gpm_email: str, actor: AppUser | None, today: date | None = None) -> tuple[str, bytes]:
    """(file name, workbook bytes) of one owner's review pack. Recorded in the audit log."""
    today = today or date.today()
    owner = owners(session).get(gpm_email.lower())
    if owner is None:
        raise ValueError(f"{gpm_email} owns no brands")
    lines = pack_lines(session, owner["brands"], today)
    last = _last_closed_month(today)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws["A1"] = f"Rebate agreements on your brands: please check - {owner['name']}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"Your brands: {', '.join(owner['brands'])}. Rates as at {today:%d %b %Y}; sales are the 12 months to "
                f"{month_end(last):%d %b %Y}, in each entity's currency.")
    ws["A3"] = ("For each line: Correct? Y or N. If N, put the right rate and say why in Comment (e.g. the retailer's "
                "email). Comment on anything else you know (e.g. a brand that should be excluded). Then save and email it back.")
    ws["A4"] = "Only the three yellow columns can be changed. Please don't add, remove or re-order rows."
    for cell in ("A2", "A3", "A4"):
        ws[cell].alignment = Alignment(wrap_text=False)
    ws.append([])
    yellow, grey = PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="E7E6E6")
    for col, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=name)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.fill = yellow if name in ANSWERS else grey
    for i, line in enumerate(lines, start=HEADER_ROW + 1):
        a = line.agreement
        values = [
            a.code, a.entity_code, a.customer_label, ", ".join(line.brands),
            line.scope_text + ("" if line.scope_confirmed else " (awaiting approval)"),
            a.rate_type.value, "Yes" if a.accrual_mode is AccrualMode.ACCRUE else "No (check only)",
            _pct(line.rate.rate) if line.rate else "none approved yet",
            line.rate.effective_from if line.rate else None, line.rate.effective_to if line.rate else None,
            (f"{_pct(line.pending.rate)} from {line.pending.effective_from:%d %b %Y}" if line.pending else ""),
            (line.rate or line.pending).source_reference,
            float(line.sales),
            None, None, None,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=i, column=col, value=value)
            if COLUMNS[col - 1] in ANSWERS:
                cell.fill = yellow
                cell.protection = Protection(locked=False)
            if isinstance(value, date):
                cell.number_format = "dd mmm yyyy"
        ws.cell(row=i, column=COLUMNS.index("Your brands' sales, last 12 months") + 1).number_format = \
            f'"{CURRENCY[a.entity_code]}"#,##0;-"{CURRENCY[a.entity_code]}"#,##0;-'
    last_row = HEADER_ROW + max(len(lines), 1)
    yes_no = DataValidation(type="list", formula1='"Y,N"', allow_blank=True, error="Y or N", showErrorMessage=True)
    ws.add_data_validation(yes_no)
    yes_no.add(f"N{HEADER_ROW + 1}:N{last_row}")
    rate = DataValidation(type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True,
                          error="A percentage between 0 and 100, e.g. 17.5", showErrorMessage=True)
    ws.add_data_validation(rate)
    rate.add(f"O{HEADER_ROW + 1}:O{last_row}")
    for col, width in zip("ABCDEFGHIJKLMNOP", (26, 7, 22, 22, 46, 10, 12, 11, 12, 12, 20, 34, 16, 10, 11, 40)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=4)
    ws.protection.sheet = True  # no password: it guards against accidents, not people
    ws.protection.formatColumns = False
    ws.protection.sort = True
    ws.protection.autoFilter = False
    ws.auto_filter.ref = f"A{HEADER_ROW}:P{last_row}"

    about = wb.create_sheet(ABOUT)
    issued = datetime.now(UTC)
    meta = {"owner": gpm_email.lower(), "issued": issued.isoformat(), "brands": owner["brands"],
            "refs": [line.agreement.code for line in lines]}
    about["A1"] = "MacGear Rebate Master - brand review pack"
    about["A1"].font = Font(bold=True)
    about["A2"] = f"For {owner['name']} ({gpm_email.lower()}), issued {issued:%d %b %Y %H:%M} UTC by Finance."
    about["A3"] = "Notes held for your brands:"
    for i, (code, note) in enumerate(sorted(owner["notes"].items()), start=4):
        about[f"A{i}"] = f"{code}: {note}"
    about["A30"] = "Please leave this cell as it is:"
    about["A31"] = json.dumps(meta)
    about.sheet_state = "visible"
    out = io.BytesIO()
    wb.save(out)
    audit(session, actor, "review_pack.issue", gpm_email.lower(), agreements=len(lines), brands=owner["brands"])
    safe = owner["name"].replace(" ", "_")
    return f"Rebate review - {safe} - {today:%Y-%m-%d}.xlsx", out.getvalue()


# ---------------------------------------------------------------- the reply


@dataclass
class ReviewResult:
    owner: str
    confirmed: list[str]
    corrections: list[tuple[str, Decimal | None, str]]  # (ref, rate the owner gives, comment)
    comments: list[tuple[str, str]]
    unanswered: list[str]


def read_return(path: Path) -> tuple[ReviewResult | None, list[str]]:
    problem = workbook_problem(path)
    if problem:
        return None, [problem]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET not in wb.sheetnames or ABOUT not in wb.sheetnames:
            return None, ["this is not a brand review pack (its sheets are missing)"]
        meta_cell = wb[ABOUT]["A31"].value
        rows = list(wb[SHEET].iter_rows(min_row=HEADER_ROW, values_only=True))
    finally:
        wb.close()
    try:
        meta = json.loads(meta_cell)
    except (TypeError, ValueError):
        return None, ["the pack's reference on the 'About this pack' sheet has been changed"]
    header = [str(h or "").strip() for h in rows[0]]
    if header[:len(COLUMNS)] != COLUMNS:
        return None, ["the columns have been changed"]
    problems = []
    result = ReviewResult(meta["owner"], [], [], [], [])
    refs = []
    for number, row in enumerate(rows[1:], start=HEADER_ROW + 1):
        ref = str(row[0] or "").strip()
        if not ref:
            continue
        refs.append(ref)
        answer, rate, comment = (row[13], row[14], str(row[15] or "").strip())
        answer = str(answer or "").strip().upper()[:1]
        given = None
        if rate not in (None, ""):
            try:
                given = (Decimal(str(rate).replace("%", "").strip()) / 100).quantize(Decimal("0.000001"))
            except InvalidOperation:
                problems.append(f"row {number} ({ref}): '{rate}' is not a rate")
                continue
        if answer == "Y" and given is None:
            result.confirmed.append(ref)
            if comment:
                result.comments.append((ref, comment))
        elif answer == "N" or given is not None:
            result.corrections.append((ref, given, comment))
        elif comment:
            result.comments.append((ref, comment))
        else:
            result.unanswered.append(ref)
    if refs != meta["refs"]:
        problems.append("rows have been added, removed or re-ordered; ask for the pack to be re-issued")
    return result, problems


def load_return(session: Session, path: Path, actor: AppUser) -> tuple[LoadBatch, ReviewResult]:
    """Record a returned pack: confirmations in the audit log; corrections and comments listed, with the
    workbook attached as evidence to each agreement it corrects or comments on."""
    path = Path(path)
    result, problems = read_return(path)
    if problems:
        raise LoadRejected(path.name, problems)
    sha = file_sha256(path)
    if session.scalar(select(LoadBatch).where(LoadBatch.kind == LoadKind.REVIEW_RETURN, LoadBatch.file_sha256 == sha)):
        raise AlreadyLoaded(f"{path.name} has already been loaded")
    agreements = {a.code: a for a in session.scalars(select(RebateAgreement))}
    unknown = [ref for ref in result.confirmed + [c[0] for c in result.corrections] if ref not in agreements]
    if unknown:
        raise LoadRejected(path.name, [f"agreement {ref} is not in the rebate master" for ref in unknown])
    batch = LoadBatch(kind=LoadKind.REVIEW_RETURN, source=f"review pack returned by {result.owner}", file_name=path.name,
                      file_sha256=sha, row_count=len(result.confirmed) + len(result.corrections) + len(result.comments),
                      loaded_by_id=actor.id, summary={})
    session.add(batch)
    session.flush()
    content = path.read_bytes()
    for ref in result.confirmed:
        audit(session, actor, "rebate.owner_confirmed", ref, owner=result.owner, batch=batch.id, file_sha256=sha)
    flagged = {ref for ref, _, _ in result.corrections} | {ref for ref, _ in result.comments}
    for ref, rate, comment in result.corrections:
        audit(session, actor, "rebate.owner_correction", ref, owner=result.owner, rate=str(rate) if rate is not None else None,
              comment=comment, batch=batch.id)
    for ref, comment in result.comments:
        audit(session, actor, "rebate.owner_comment", ref, owner=result.owner, comment=comment, batch=batch.id)
    for ref in sorted(flagged):
        attach_evidence(session, actor=actor, agreement=agreements[ref],
                        uploads=[Upload(path.name, content, f"Brand review returned by {result.owner}")])
    batch.summary = {
        "owner": result.owner, "confirmed": len(result.confirmed), "corrections": len(result.corrections),
        "comments": len(result.comments), "unanswered": len(result.unanswered),
        "notes": [f"{ref}: says {_pct(rate) if rate is not None else 'not correct'}"
                  + (f" - {comment}" if comment else "") + " (enter any change as a proposal)"
                  for ref, rate, comment in result.corrections]
                 + [f"{ref}: comment - {comment}" for ref, comment in result.comments],
    }
    session.flush()
    return batch, result
